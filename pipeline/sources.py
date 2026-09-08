"""
Non-US sanctions lists, normalised to the same row shape build.py uses for the
US Consolidated Screening List:

    {authority, source, entity_number, type, programs, name, addresses, remarks,
     alt_names, nationalities, dates_of_birth, places_of_birth, source_information_url, ids}

Every fetcher is wrapped so a broken download or a changed format skips that
authority and reports it in meta.json instead of failing the whole build.
Add --sample-dir to build.py to test with local copies of the files.
"""
import csv, io, re, os, zipfile
import xml.etree.ElementTree as ET

UA = {"User-Agent": "Mozilla/5.0 (compatible; sanctionscope-build; +https://github.com/)"}
SOURCES = {
    "EU": ("EU Consolidated Financial Sanctions List",
           "https://webgate.ec.europa.eu/fsd/fsf/public/files/csvFullSanctionsList_1_1/content?token=dG9rZW4tMjAxNw", "eu.csv"),
    "UK": ("UK Sanctions List (FCDO)", "https://sanctionslist.fcdo.gov.uk/docs/UK-Sanctions-List.csv", "uk.csv"),
    "UN": ("UN Security Council Consolidated List", "https://scsanctions.un.org/resources/xml/en/consolidated.xml", "un.xml"),
    "AU": ("Australia DFAT Consolidated List", "https://www.dfat.gov.au/sites/default/files/Australian_Sanctions_Consolidated_List.xlsx", "au.xlsx"),
    "CA": ("Canada SEMA / Autonomous Sanctions", "https://www.international.gc.ca/world-monde/assets/office_docs/international_relations-relations_internationales/sanctions/sema-lmes.xml", "ca.xml"),
}
RECORD_URL = {
    "EU": "https://www.sanctionsmap.eu/#/main", "UK": "https://www.gov.uk/government/publications/the-uk-sanctions-list",
    "UN": "https://main.un.org/securitycouncil/en/content/un-sc-consolidated-list", "AU": "https://www.dfat.gov.au/international-relations/security/sanctions/consolidated-list",
    "CA": "https://www.international.gc.ca/world-monde/international_relations-relations_internationales/sanctions/consolidated-consolide.aspx",
}

def fetch(auth, sample_dir=None):
    name, url, fname = SOURCES[auth]
    if sample_dir:
        p = os.path.join(sample_dir, fname)
        if not os.path.exists(p): raise FileNotFoundError(p)
        with open(p, "rb") as f: return f.read()
    import requests
    r = requests.get(url, timeout=240, headers=UA)
    r.raise_for_status()
    return r.content

def clean(s):
    return re.sub(r"\s+", " ", (s or "")).strip()

def joinaddr(*parts):
    return ", ".join(clean(x) for x in parts if clean(x))

def row(auth, eid, name, typ, programs, addresses, remarks="", alt=(), nat=(), dob="", pob="", ids="", url="", listed=""):
    return {"authority": auth, "source": SOURCES[auth][0], "entity_number": f"{auth}-{eid}", "type": typ,
            "programs": "; ".join(sorted({f"{auth}:{p}" for p in programs if p})),
            "name": clean(name), "addresses": "; ".join(dict.fromkeys(a for a in addresses if a)),
            "remarks": clean(remarks), "alt_names": "; ".join(dict.fromkeys(clean(a) for a in alt if clean(a) and clean(a).lower() != clean(name).lower())),
            "nationalities": "; ".join(dict.fromkeys(clean(n) for n in nat if clean(n))), "dates_of_birth": clean(dob), "places_of_birth": clean(pob),
            "ids": clean(ids), "source_information_url": url or RECORD_URL[auth], "vessel_flag": "", "start_date": clean(listed)}

# ------------------------------------------------------------------ EU
def parse_eu(raw):
    text = raw.decode("utf-8-sig", errors="replace")
    rd = csv.DictReader(io.StringIO(text), delimiter=";")
    groups = {}
    for r in rd:
        g = r.get("Entity_LogicalId") or r.get("Entity_EU_ReferenceNumber")
        if not g: continue
        e = groups.setdefault(g, {"names": [], "addr": [], "nat": [], "dob": [], "pob": [], "prog": set(), "type": "", "remark": "", "ref": "", "ids": [], "url": "", "date": ""})
        st = (r.get("Entity_SubjectType") or "").lower()
        if st: e["type"] = "Individual" if st.startswith("p") else "Entity"
        e["prog"].add(clean(r.get("Entity_Regulation_Programme")))
        e["remark"] = e["remark"] or clean(r.get("Entity_Remark"))
        e["ref"] = e["ref"] or clean(r.get("Entity_EU_ReferenceNumber"))
        e["url"] = e["url"] or clean(r.get("Entity_Regulation_PublicationUrl"))
        e["date"] = e["date"] or clean(r.get("Entity_DesignationDate")) or clean(r.get("Entity_Regulation_EntryIntoForceDate"))
        wn = clean(r.get("NameAlias_WholeName")) or joinaddr(r.get("NameAlias_FirstName"), r.get("NameAlias_MiddleName"), r.get("NameAlias_LastName")).replace(",", "")
        if wn and wn not in e["names"]: e["names"].append(wn)
        a = joinaddr(r.get("Address_Street"), r.get("Address_City"), r.get("Address_Region"), r.get("Address_CountryDescription") or r.get("Address_CountryIso2Code"))
        if a and a not in e["addr"]: e["addr"].append(a)
        n = clean(r.get("Citizenship_CountryDescription") or r.get("Citizenship_CountryIso2Code"))
        if n and n not in e["nat"]: e["nat"].append(n)
        d = clean(r.get("BirthDate_BirthDate")) or clean(r.get("BirthDate_Year"))
        if d and d not in e["dob"]: e["dob"].append(d)
        pb = joinaddr(r.get("BirthDate_City"), r.get("BirthDate_CountryDescription"))
        if pb and pb not in e["pob"]: e["pob"].append(pb)
        idn = clean(r.get("Identification_Number"))
        if idn and idn not in e["ids"]: e["ids"].append(f"{clean(r.get('Identification_TypeDescription')) or 'ID'} {idn}")
    out = []
    for g, e in groups.items():
        if not e["names"]: continue
        out.append(row("EU", g, e["names"][0], e["type"] or "Entity", e["prog"], e["addr"], e["remark"], e["names"][1:], e["nat"],
                       "; ".join(e["dob"]), "; ".join(e["pob"]), "; ".join(e["ids"]), e["url"], e["date"]))
    return out

# ------------------------------------------------------------------ UK
UK_COLS = {   # canonical -> accepted header names (UKSL 2025 format first, retired OFSI ConList second)
    "id": ["Unique ID", "Group ID"], "n6": ["Name 6"], "n1": ["Name 1"], "n2": ["Name 2"], "n3": ["Name 3"], "n4": ["Name 4"], "n5": ["Name 5"],
    "ntype": ["Name type", "Name Type", "Alias Type"], "regime": ["Regime Name", "Regime"], "kind": ["Individual, Entity, Ship", "Group Type"],
    "info": ["Other Information"], "reasons": ["UK Statement of Reasons"], "a1": ["Address Line 1", "Address 1"], "a2": ["Address Line 2", "Address 2"],
    "a3": ["Address Line 3", "Address 3"], "a4": ["Address Line 4", "Address 4"], "a5": ["Address Line 5", "Address 5"], "a6": ["Address Line 6", "Address 6"],
    "post": ["Address Postal Code", "Post/Zip Code"], "country": ["Address Country", "Country"], "date": ["Date Designated", "UK Sanctions List Date Designated", "Listed On"],
    "dob": ["D.O.B", "DOB"], "nat": ["Nationality(/ies)", "Nationality"], "nid": ["National Identifier number", "National Identification Number"],
    "pass": ["Passport number", "Passport Number"], "pos": ["Position"], "tob": ["Town of birth", "Town of Birth"], "cob": ["Country of birth", "Country of Birth"],
    "flag": ["Current believed flag of ship"], "imo": ["IMO number"], "parent": ["Parent Company"], "subs": ["Subsidiaries"],
}
def parse_uk(raw):
    text = raw.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    start = next((i for i, l in enumerate(lines[:10]) if ("Unique ID" in l or "Group ID" in l) and "Name 6" in l), 0)
    rd = csv.DictReader(io.StringIO("\n".join(lines[start:])))
    hdr = {clean(h).lower(): h for h in (rd.fieldnames or [])}
    col = {}
    for k, names in UK_COLS.items():
        for nm in names:
            if nm.lower() in hdr: col[k] = hdr[nm.lower()]; break
    if "id" not in col or "n6" not in col: raise ValueError("UK list: expected columns not found; header was " + ", ".join(list(hdr)[:12]))
    g = lambda r, k: clean(r.get(col[k])) if k in col else ""
    groups = {}
    for r in rd:
        gid = g(r, "id")
        if not gid: continue
        e = groups.setdefault(gid, {"primary": "", "alias": [], "addr": [], "nat": [], "dob": [], "pob": [], "prog": set(), "type": "", "info": "", "pos": "", "ids": [], "date": "", "flag": ""})
        nm = joinaddr(g(r, "n6"), g(r, "n1"), g(r, "n2"), g(r, "n3"), g(r, "n4"), g(r, "n5")).replace(",", "") if g(r, "kind").lower().startswith("ind") \
             else joinaddr(g(r, "n1"), g(r, "n2"), g(r, "n3"), g(r, "n4"), g(r, "n5"), g(r, "n6")).replace(",", "")
        # individuals: surname is Name 6, so put it last for a natural order
        if g(r, "kind").lower().startswith("ind"):
            nm = joinaddr(g(r, "n1"), g(r, "n2"), g(r, "n3"), g(r, "n4"), g(r, "n5"), g(r, "n6")).replace(",", "")
        nt = g(r, "ntype").lower()
        if nt.startswith("primary name") and "variation" not in nt and not e["primary"]: e["primary"] = nm
        elif nm and nm not in e["alias"]: e["alias"].append(nm)
        gt = g(r, "kind").lower()
        e["type"] = "Individual" if gt.startswith("ind") else "Vessel" if gt.startswith("ship") else "Entity"
        e["prog"].add(g(r, "regime"))
        e["info"] = e["info"] or g(r, "info")
        e["pos"] = e["pos"] or g(r, "pos")
        e["date"] = e["date"] or g(r, "date")
        e["flag"] = e["flag"] or g(r, "flag")
        a = joinaddr(g(r, "a1"), g(r, "a2"), g(r, "a3"), g(r, "a4"), g(r, "a5"), g(r, "a6"), g(r, "post"), g(r, "country"))
        if a and a not in e["addr"]: e["addr"].append(a)
        for n in re.split(r"[,;/]|\(\d+\)", g(r, "nat")):
            n = clean(n)
            if n and n not in e["nat"]: e["nat"].append(n)
        d = g(r, "dob")
        if d and d not in e["dob"]: e["dob"].append(d)
        pb = joinaddr(g(r, "tob"), g(r, "cob"))
        if pb and pb not in e["pob"]: e["pob"].append(pb)
        for k, lab in (("pass", "Passport"), ("nid", "National ID"), ("imo", "IMO")):
            v = g(r, k)
            if v and f"{lab} {v}" not in e["ids"]: e["ids"].append(f"{lab} {v}")
        for k, lab in (("parent", "Parent"), ("subs", "Subsidiaries")):
            v = g(r, k)
            if v and lab not in e["info"]: e["info"] = "; ".join(x for x in [e["info"], f"{lab}: {v}"] if x)
    out = []
    for gid, e in groups.items():
        name = e["primary"] or (e["alias"][0] if e["alias"] else "")
        if not name: continue
        alias = [a for a in e["alias"] if a != name]
        rem = "; ".join(x for x in [e["pos"], e["info"]] if x)
        r0 = row("UK", gid, name, e["type"], e["prog"], e["addr"], rem, alias, e["nat"], "; ".join(e["dob"]), "; ".join(e["pob"]), "; ".join(e["ids"]),
                 f"https://search-uk-sanctions-list.service.gov.uk/designations/{gid}", e["date"])
        r0["vessel_flag"] = e["flag"]
        out.append(r0)
    return out

# ------------------------------------------------------------------ UN
def parse_un(raw):
    root = ET.fromstring(raw)
    def tx(el, tag):
        x = el.find(tag)
        return clean(x.text) if x is not None and x.text else ""
    out = []
    for ind in root.iter("INDIVIDUAL"):
        name = " ".join(x for x in [tx(ind, "FIRST_NAME"), tx(ind, "SECOND_NAME"), tx(ind, "THIRD_NAME"), tx(ind, "FOURTH_NAME")] if x)
        if not name: continue
        alias = [tx(a, "ALIAS_NAME") for a in ind.findall("INDIVIDUAL_ALIAS")]
        addr = [joinaddr(tx(a, "STREET"), tx(a, "CITY"), tx(a, "STATE_PROVINCE"), tx(a, "COUNTRY")) for a in ind.findall("INDIVIDUAL_ADDRESS")]
        nat = [clean(v.text) for n in ind.findall("NATIONALITY") for v in n.findall("VALUE") if v.text]
        dob = [tx(d, "DATE") or tx(d, "YEAR") for d in ind.findall("INDIVIDUAL_DATE_OF_BIRTH")]
        pob = [joinaddr(tx(p, "CITY"), tx(p, "STATE_PROVINCE"), tx(p, "COUNTRY")) for p in ind.findall("INDIVIDUAL_PLACE_OF_BIRTH")]
        ids = [joinaddr(tx(d, "TYPE_OF_DOCUMENT"), tx(d, "NUMBER"), tx(d, "ISSUING_COUNTRY")) for d in ind.findall("INDIVIDUAL_DOCUMENT")]
        rem = "; ".join(x for x in [tx(ind, "DESIGNATION/VALUE"), tx(ind, "COMMENTS1")] if x)
        out.append(row("UN", tx(ind, "DATAID"), name, "Individual", [tx(ind, "UN_LIST_TYPE")], addr, rem, alias, nat, "; ".join(dob), "; ".join(pob), "; ".join(ids),
                       "https://main.un.org/securitycouncil/en/content/un-sc-consolidated-list", tx(ind, "LISTED_ON")))
    for ent in root.iter("ENTITY"):
        name = tx(ent, "FIRST_NAME")
        if not name: continue
        alias = [tx(a, "ALIAS_NAME") for a in ent.findall("ENTITY_ALIAS")]
        addr = [joinaddr(tx(a, "STREET"), tx(a, "CITY"), tx(a, "STATE_PROVINCE"), tx(a, "COUNTRY")) for a in ent.findall("ENTITY_ADDRESS")]
        rem = tx(ent, "COMMENTS1")
        out.append(row("UN", tx(ent, "DATAID"), name, "Entity", [tx(ent, "UN_LIST_TYPE")], addr, rem, alias, [], "", "", "",
                       "https://main.un.org/securitycouncil/en/content/un-sc-consolidated-list", tx(ent, "LISTED_ON")))
    return out

# ------------------------------------------------------------------ AU
def parse_au(raw):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = None
    groups = {}
    for r in rows:
        vals = ["" if v is None else str(v) for v in r]
        if header is None:
            if any("Name of Individual" in v for v in vals): header = [clean(v) for v in vals]
            continue
        d = dict(zip(header, vals))
        ref = re.sub(r"[a-zA-Z]+$", "", clean(d.get("Reference")).split(".")[0])   # "1000a" -> "1000"
        if not ref: continue
        e = groups.setdefault(ref, {"primary": "", "alias": [], "addr": [], "nat": [], "dob": [], "pob": [], "prog": set(), "type": "", "info": "", "date": ""})
        nm = clean(d.get("Name of Individual or Entity"))
        ntype = (d.get("Name Type") or "").lower()
        if ntype.startswith("primary") and not e["primary"]: e["primary"] = nm
        elif nm and "script" not in ntype: e["alias"].append(nm)
        elif nm: e["alias"].append(nm)
        t = (d.get("Type") or "").lower(); e["type"] = "Individual" if t.startswith(("ind", "per")) else "Vessel" if t.startswith("ves") else "Entity"
        for c in re.split(r"[;/]", d.get("Committees") or d.get("Listing Information") or ""):
            c = clean(c)
            if c: e["prog"].add(c[:40])
        a = clean(d.get("Address"))
        if a and a not in e["addr"]: e["addr"].append(a)
        for n in re.split(r"[,;/]", d.get("Citizenship") or ""):
            n = clean(n)
            if n and n not in e["nat"]: e["nat"].append(n)
        dob = clean(d.get("Date of Birth"))
        if dob and dob not in e["dob"]: e["dob"].append(dob)
        pb = clean(d.get("Place of Birth"))
        if pb and pb not in e["pob"]: e["pob"].append(pb)
        e["info"] = e["info"] or clean(d.get("Additional Information"))
        e["date"] = e["date"] or clean(d.get("Control Date"))
    out = []
    for ref, e in groups.items():
        name = e["primary"] or (e["alias"][0] if e["alias"] else "")
        if not name: continue
        out.append(row("AU", ref, name, e["type"], e["prog"] or {"Autonomous"}, e["addr"], e["info"], [a for a in e["alias"] if a != name], e["nat"], "; ".join(e["dob"]), "; ".join(e["pob"]), "", "", e["date"]))
    return out

# ------------------------------------------------------------------ CA
def parse_ca(raw):
    """Tags are bilingual, e.g. <Country-Pays>, <LastName-NomDeFamille>, <GivenName-Prenom>, <Entity-Entite>,
    <Aliases-Alias>, <Schedule-Annexe>, <Item-NumeroDarticle>, <DateOfListing-DateDinscription>. Match on the English prefix."""
    root = ET.fromstring(raw)
    def tx(el, prefix):
        for ch in el:
            tag = ch.tag.split("}")[-1]
            if tag.lower().startswith(prefix.lower()):
                return clean(ch.text or "")
        return ""
    out = []
    i = 0
    for rec in root.iter("record"):
        i += 1
        country = tx(rec, "Country").split("/")[0].strip()
        ent = tx(rec, "Entity"); gn = tx(rec, "GivenName"); ln = tx(rec, "LastName")
        kind = tx(rec, "EntityOrShip") or tx(rec, "Type")
        if ent:
            name, typ = ent, ("Vessel" if "ship" in kind.lower() or tx(rec, "ShipIMO") else "Entity")
        elif gn or ln:
            name, typ = f"{gn} {ln}".strip(), "Individual"
        else:
            continue
        alias = [clean(a) for a in re.split(r";|\n", tx(rec, "Aliases")) if clean(a)]
        prog = country or "Autonomous"
        sched = tx(rec, "Schedule"); item = tx(rec, "Item")
        rem = "; ".join(x for x in [tx(rec, "Title"), sched and f"Schedule {sched}", item and f"Item {item}"] if x)
        r0 = row("CA", f"{prog}-{sched}-{item}".replace(" ", "") if item else i, name, typ, [prog], [country] if typ != "Individual" and country else [],
                 rem, alias, [country] if typ == "Individual" and country else [], tx(rec, "DateOfBirth"), "", tx(rec, "ShipIMO") and "IMO " + tx(rec, "ShipIMO"), "",
                 tx(rec, "DateOfListing"))
        out.append(r0)
    return out

PARSERS = {"EU": parse_eu, "UK": parse_uk, "UN": parse_un, "AU": parse_au, "CA": parse_ca}

def load_all(sample_dir=None, only=None):
    """Returns (rows, status) where status[auth] = {'ok': bool, 'n': int, 'error': str}."""
    rows, status = [], {}
    for auth in SOURCES:
        if only and auth not in only: continue
        try:
            raw = fetch(auth, sample_dir)
            parsed = PARSERS[auth](raw)
            if len(parsed) < 50: raise ValueError(f"only {len(parsed)} records parsed; format probably changed")
            rows.extend(parsed); status[auth] = {"ok": True, "n": len(parsed), "error": ""}
        except Exception as ex:  # a broken source must not sink the build
            status[auth] = {"ok": False, "n": 0, "error": f"{type(ex).__name__}: {str(ex)[:200]}"}
    return rows, status
